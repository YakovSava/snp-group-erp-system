import io

import magic

IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


def sniff_mime(data: bytes) -> str:
    return magic.from_buffer(data[:4096], mime=True)


def is_image(mime: str) -> bool:
    return mime in IMAGE_MIMES


def extract_text(data: bytes, mime: str) -> str | None:
    """Best-effort text extraction for files attached to a chat message, so
    the model can answer questions about documents it can't natively read.
    Returns None for images (handled as vision input instead) or formats we
    don't know how to parse.
    """
    try:
        if mime == "application/pdf":
            return _extract_pdf(data)
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _extract_docx(data)
        if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return _extract_xlsx(data)
        if mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return _extract_pptx(data)
        if mime.startswith("text/"):
            return data.decode("utf-8", errors="replace")
    except Exception:
        return None
    return None


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs if p.text)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell) for cell in row if cell is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    lines = []
    for i, slide in enumerate(presentation.slides, start=1):
        lines.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)
